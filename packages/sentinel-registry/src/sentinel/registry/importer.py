"""Bulk camera onboarding from CSV and Excel.

Fifty cameras is a spreadsheet, eighty thousand is a spreadsheet too. The
import is row-at-a-time and partial: every row that validates is written,
every row that does not comes back with a reason. An import that fails
wholesale on row 400 of 900 is useless to whoever is doing the onboarding.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date
from typing import Any

import asyncpg
from sentinel.core.logging import get_logger
from sentinel.core.models import CameraIn
from sentinel.core.types import CameraKind
from sentinel.registry import repo

log = get_logger(__name__)

#: Spreadsheet header -> CameraIn field. Anything else in the file is ignored,
#: so the departments can keep their own extra columns.
COLUMN_MAP = {
    "camera_id": "camera_id",
    "id": "camera_id",
    "department": "department_code",
    "department_code": "department_code",
    "name": "name",
    "kind": "kind",
    "type": "kind",
    "vendor": "vendor",
    "protocol": "protocol",
    "lat": "lat",
    "latitude": "lat",
    "lon": "lon",
    "lng": "lon",
    "longitude": "lon",
    "bearing": "bearing_deg",
    "bearing_deg": "bearing_deg",
    "range_m": "range_m",
    "range": "range_m",
    "storage": "storage_kind",
    "storage_kind": "storage_kind",
    "retention_days": "retention_days",
    "retention": "retention_days",
    "contract_expiry": "contract_expiry",
}

REQUIRED = {"camera_id", "name", "kind", "lat", "lon"}


@dataclass
class RowError:
    row: int
    camera_id: str | None
    reason: str


@dataclass
class ImportResult:
    accepted: int = 0
    rejected: list[RowError] = field(default_factory=list)
    camera_ids: list[str] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        return {
            "accepted": self.accepted,
            "rejected": [vars(e) for e in self.rejected],
            "camera_ids": self.camera_ids,
        }


def _normalise(raw: dict[str, Any]) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for key, value in raw.items():
        if key is None:
            continue
        field_name = COLUMN_MAP.get(str(key).strip().lower())
        if field_name and value not in (None, ""):
            out[field_name] = value
    return out


def _coerce(row: dict[str, Any]) -> dict[str, Any]:
    out = dict(row)
    for numeric in ("lat", "lon", "bearing_deg", "range_m"):
        if numeric in out:
            out[numeric] = float(out[numeric])
    if "retention_days" in out:
        out["retention_days"] = int(float(out["retention_days"]))
    if "contract_expiry" in out and not isinstance(out["contract_expiry"], date):
        out["contract_expiry"] = date.fromisoformat(str(out["contract_expiry"])[:10])
    if "kind" in out:
        k = str(out["kind"]).strip().lower()
        out["kind"] = CameraKind.ANALOG if k.startswith("a") else CameraKind.IP
    return out


def parse_csv(data: bytes) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig", errors="replace")
    return [_normalise(r) for r in csv.DictReader(io.StringIO(text))]


def parse_xlsx(data: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = wb[wb.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    try:
        header = [str(h).strip().lower() if h is not None else None for h in next(rows)]
    except StopIteration:
        return []
    return [_normalise(dict(zip(header, values, strict=False))) for values in rows]


async def import_rows(
    conn: asyncpg.Connection, rows: list[dict[str, Any]], default_department_id: int | None
) -> ImportResult:
    """Validate and upsert. Each row stands or falls on its own."""
    result = ImportResult()
    dept_cache: dict[str, int] = {}

    for index, raw in enumerate(rows, start=2):   # row 1 is the header
        camera_id = str(raw.get("camera_id", "")).strip() or None
        try:
            missing = REQUIRED - raw.keys()
            if missing:
                raise ValueError(f"missing required column(s): {', '.join(sorted(missing))}")

            data = _coerce(raw)
            code = str(data.pop("department_code", "") or "").strip().upper()
            if code:
                if code not in dept_cache:
                    dept = await repo.get_department_by_code(conn, code)
                    if dept is None or dept.id is None:
                        raise ValueError(f"unknown department code {code!r}")
                    dept_cache[code] = dept.id
                data["department_id"] = dept_cache[code]
            elif default_department_id is not None:
                data["department_id"] = default_department_id
            else:
                raise ValueError("no department column and no default department given")

            camera = CameraIn(**data)
            await repo.upsert_camera(conn, camera)
            result.accepted += 1
            result.camera_ids.append(camera.camera_id)
        except Exception as exc:
            result.rejected.append(RowError(row=index, camera_id=camera_id, reason=str(exc)))

    log.info("bulk_import", accepted=result.accepted, rejected=len(result.rejected))
    return result
